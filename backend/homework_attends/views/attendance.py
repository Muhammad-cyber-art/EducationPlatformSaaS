from django.db import models
from django.db.models import Q
from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.permissions import IsAuthenticated,AllowAny
from django.shortcuts import get_object_or_404
from django.utils import timezone
from django.http import HttpResponse

import logging
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import timedelta

from homework_attends.models import Attendance
from groups.models import Group
from homework_attends.serializers import AttendanceSerializer, AttendanceEditSerializer
from homework_attends.services import (
    get_or_create_attendance_records, 
    generate_weekly_attendance_report, 
    bulk_confirm_attendance, 
    get_monthly_attendance_data,
    edit_past_attendance,
    get_last_3_lesson_dates
)
from permissions.permissions import HasModulePermission
from reports.models import ReportDownloadTrack
from datetime import date as date_type

logger = logging.getLogger(__name__)

class AttendanceViewSet(viewsets.ModelViewSet):
    serializer_class = AttendanceSerializer
    permission_classes = [AllowAny]
    pagination_class = None
    module_name = 'homework'
    http_method_names = ['get', 'post', 'patch', 'delete', 'head', 'options']

    def get_permissions(self):
        """Monthly report va student_monthly uchun autentifikatsiya majburiy"""
        if self.action in ['monthly_report', 'student_monthly']:
            return [IsAuthenticated()]
        return super().get_permissions()

    def get_queryset(self):
       group_id = self.request.query_params.get('group_id')
       date_str = self.request.query_params.get('date', str(timezone.localdate()))
       
       if not group_id:
           return Attendance.objects.none()
       
       try:
           requested_date = timezone.datetime.strptime(date_str, '%Y-%m-%d').date()
       except (ValueError, TypeError):
           requested_date = timezone.localdate()

       group = get_object_or_404(Group, id=group_id)
       today = timezone.localdate()
       last_3_dates = get_last_3_lesson_dates(group)
       is_past_month = requested_date.year < today.year or (requested_date.year == today.year and requested_date.month < today.month)
       has_full_attendance = self._user_has_full_attendance(self.request.user)
       within_1_month = has_full_attendance and (today - timedelta(days=30) <= requested_date <= today)
       view_only = is_past_month and (requested_date not in last_3_dates) and not within_1_month
       return get_or_create_attendance_records(group, requested_date, view_only=view_only)

    @action(detail=False, methods=['get'])
    def weekly_report(self, request):
         group_id = request.query_params.get('group_id')
         group = get_object_or_404(Group, id=group_id)
         
         report = generate_weekly_attendance_report(group)
         return Response(report)
    
    def create(self, request, *args, **kwargs):
        attendance_id = request.data.get('id')
        is_present = request.data.get('is_present')
        student_id = request.data.get('student_id')
        date_str = request.data.get('date')

        # 1. Update existing record
        if attendance_id is not None:
            instance = get_object_or_404(Attendance, id=attendance_id)
            if request.user.is_authenticated and not self._user_can_access_group(request.user, instance.group):
                return Response({"detail": "Sizda ushbu guruh davomatini tahrirlash huquqi yo'q"}, status=status.HTTP_403_FORBIDDEN)

            today = timezone.localdate()
            if instance.date != today:
                last_3_dates = get_last_3_lesson_dates(instance.group)
                
                # Grace period: Tungi 4 gacha kechagi kunni o'zgartirishga ruxsat
                is_early_morning = timezone.localtime().hour < 4
                is_yesterday = instance.date == (today - timedelta(days=1))
                has_full_attendance = self._user_has_full_attendance(request.user)
                within_1_month = has_full_attendance and (today - timedelta(days=30) <= instance.date <= today)
                if instance.date not in last_3_dates and not (is_early_morning and is_yesterday) and not within_1_month:
                    return Response({"detail": f"Faqat bugungi yoki oxirgi 3 ta o'tilgan dars davomatini tahrirlashingiz mumkin."}, status=status.HTTP_400_BAD_REQUEST)
            
            instance.is_present = is_present
            instance.marked_by = request.user
            instance.save()

            # Moliya va maoshni qayta hisoblash
            try:
                from finance.services import update_attendance_based_payments
                update_attendance_based_payments(instance.student, instance.group, instance.date)
            except Exception as e:
                logger.error("Failed to update payments on attendance update: %s", e)

            serializer = self.get_serializer(instance)
            return Response(serializer.data, status=status.HTTP_200_OK)

        # 2. Create new record
        group_id = request.data.get('group_id')
        if not (student_id and date_str and group_id):
             return Response({"detail": "Student ID, group_id va kun (date) majburiy"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            requested_date = timezone.datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            return Response({"detail": "Sana formati noto'g'ri (YYYY-MM-DD)"}, status=status.HTTP_400_BAD_REQUEST)

        group = get_object_or_404(Group, id=group_id)
        if request.user.is_authenticated and not self._user_can_access_group(request.user, group):
            return Response({"detail": "Sizda ushbu guruh davomatini tahrirlash huquqi yo'q"}, status=status.HTTP_403_FORBIDDEN)
        
        # Ruxsat etilgan sanalar: Bugun, yoki oxirgi 3 ta dars, yoki to'liq davomat huquqi bo'lsa 1 oy
        today = timezone.localdate()
        if requested_date != today:
            last_3_dates = get_last_3_lesson_dates(group)
            
            is_early_morning = timezone.localtime().hour < 4
            is_yesterday = requested_date == (today - timedelta(days=1))
            has_full_attendance = self._user_has_full_attendance(request.user)
            within_1_month = has_full_attendance and (today - timedelta(days=30) <= requested_date <= today)
            if requested_date not in last_3_dates and not (is_early_morning and is_yesterday) and not within_1_month:
                return Response({"detail": f"Faqat bugungi yoki oxirgi 3 ta o'tilgan dars davomatini tahrirlashingiz mumkin."}, status=status.HTTP_400_BAD_REQUEST)

        student = get_object_or_404(group.students, id=student_id)

        # Duplicate check
        attendance, created = Attendance.objects.get_or_create(
            student=student, 
            group=group, 
            date=requested_date,
            defaults={
                'is_present': is_present,
                'marked_by': request.user
            }
        )
        if not created:
            attendance.is_present = is_present
            attendance.marked_by = request.user
            attendance.save()

        # Moliya va maoshni qayta hisoblash
        try:
            from finance.services import update_attendance_based_payments
            update_attendance_based_payments(student, group, requested_date)
        except Exception as e:
            logger.error("Failed to update payments on attendance create: %s", e)

        serializer = self.get_serializer(attendance)
        return Response(serializer.data, status=status.HTTP_201_CREATED)

    @action(detail=False, methods=['post'], url_path='confirm')
    def confirm_attendance(self, request):
        group_id = request.data.get('group_id')
        date_str = request.data.get('date')
        attendances_payload = request.data.get('attendances', [])

        if not group_id or not date_str:
            return Response({"detail": "group_id va date majburiy"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            requested_date = timezone.datetime.strptime(date_str, '%Y-%m-%d').date()
        except ValueError:
            return Response({"detail": "Sana formati noto'g'ri (YYYY-MM-DD)"}, status=status.HTTP_400_BAD_REQUEST)

        group = get_object_or_404(Group, id=group_id)
        if request.user.is_authenticated and not self._user_can_access_group(request.user, group):
            return Response({"detail": "Sizda ushbu guruh davomatini tasdiqlash huquqi yo'q"}, status=status.HTTP_403_FORBIDDEN)
        
        # Ruxsat etilgan sanalar: Bugun, yoki oxirgi 3 ta dars, yoki to'liq davomat huquqi bo'lsa 1 oy
        today = timezone.localdate()
        if requested_date != today:
            last_3_dates = get_last_3_lesson_dates(group)
            
            is_early_morning = timezone.localtime().hour < 4
            is_yesterday = requested_date == (today - timedelta(days=1))
            has_full_attendance = self._user_has_full_attendance(request.user)
            within_1_month = has_full_attendance and (today - timedelta(days=30) <= requested_date <= today)
            if requested_date not in last_3_dates and not (is_early_morning and is_yesterday) and not within_1_month:
                return Response({"detail": f"Faqat bugungi yoki oxirgi 3 ta o'tilgan dars davomatini tahrirlashingiz mumkin."}, status=status.HTTP_400_BAD_REQUEST)


        try:
            queryset = bulk_confirm_attendance(group, requested_date, attendances_payload, request.user)
            serializer = AttendanceSerializer(queryset, many=True)
            return Response(serializer.data, status=status.HTTP_200_OK)
        except Exception as e:
            logger.exception(f"Attendance confirm failed for group {group_id}: {e}")
            return Response({"detail": f"Davomatni tasdiqlashda xatolik: {str(e)}"}, status=status.HTTP_400_BAD_REQUEST)

    def _user_has_full_attendance(self, user):
        """Admin yoki super admin uchun to'liq davomat (1 oylik tahrirlash) huquqini tekshirish"""
        if not user or not getattr(user, 'is_authenticated', False):
            return False
        if getattr(user, 'role', '') == 'super_admin' or getattr(user, 'is_superuser', False):
            return True
        if getattr(user, 'role', '') == 'admin':
            try:
                from permissions.models import StaffPermission
                sp = StaffPermission.objects.filter(user=user).first()
                if sp and sp.permissions:
                    return bool(sp.permissions.get('full_attendance'))
            except Exception:
                return False
        return False

    def _user_can_access_group(self, user, group):
        """Foydalanuvchi guruhga kirish huquqini tekshirish"""
        if not user or not getattr(user, 'is_authenticated', False):
            return False
        if user.role == 'super_admin':
            return True
        if user.role == 'admin':
            if self._user_has_full_attendance(user):
                return True
            allowed = [user.branch_id] if user.branch_id else []
            if hasattr(user, 'branch_accesses'):
                allowed.extend(user.branch_accesses.values_list('branch_id', flat=True))
            return group.branch_id in allowed
        if user.role == 'mentor':
            return group.mentor == user or group.additional_mentors.filter(mentor=user).exists()
        return False

    @action(detail=False, methods=['get'], url_path='monthly-report')
    def monthly_report(self, request):
        group_id = request.query_params.get('group_id')
        if not group_id:
            return Response({"detail": "group_id majburiy"}, status=status.HTTP_400_BAD_REQUEST)
        
        try:
            group = get_object_or_404(Group, id=group_id)
            
            # Foydalanuvchining guruhga kirish huquqini tekshirish
            if not self._user_can_access_group(request.user, group):
                logger.warning(
                    "User %s (role=%s) has no access to group %s (branch=%s)",
                    request.user, request.user.role, group.id, group.branch_id
                )
                return Response(
                    {"detail": "Sizda ushbu guruh davomatini ko'rish huquqi yo'q"},
                    status=status.HTTP_403_FORBIDDEN
                )
            
            today = timezone.localdate()
            
            try:
                month = int(request.query_params.get('month', today.month))
                year = int(request.query_params.get('year', today.year))
                # Oy va yil qiymatlarini tekshirish
                date_type(year, month, 1)
            except (ValueError, TypeError):
                month, year = today.month, today.year
                
            date_list, students, att_data = get_monthly_attendance_data(group, month, year)
        except Exception as e:
            logger.exception("Monthly report data olishda xatolik: group=%s, month=%s-%s", group_id, year, month)
            return Response(
                {"detail": f"Hisobot ma'lumotlarini olishda xatolik: {str(e)}"},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )
        
        export_mode = request.query_params.get('export', 'json')
        
        if export_mode == 'excel':
            wb = openpyxl.Workbook()
            ws = wb.active
            
            # openpyxl sheet title has strict rules: max 31 chars, no forbidden chars: \ / ? * : [ ]
            safe_sheet_name = f"Davomat_{group.name[:20]}"
            for char in r"\/?:*[]":
                safe_sheet_name = safe_sheet_name.replace(char, "")
            safe_sheet_name = safe_sheet_name.strip()[:30]
            if not safe_sheet_name:
                safe_sheet_name = "Davomat"
            ws.title = safe_sheet_name
            
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            center_align = Alignment(horizontal='center', vertical='center')
            
            headers = ["№", "O'quvchi ismi-familiyasi", "Telefon", "Ota-ona telefon"] + [d.day for d in date_list]
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=2, column=col_num)
                cell.value = header

            ws.column_dimensions['A'].width = 5
            ws.column_dimensions['B'].width = 30
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 15
            for col_num in range(5, len(headers) + 1):
                ws.column_dimensions[get_column_letter(col_num)].width = 4
                
            header_row = ws[2]
            for cell in header_row:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align

            frozen_fill = PatternFill(start_color="BAE6FD", end_color="BAE6FD", fill_type="solid")
            frozen_font = Font(color="0369A1", bold=True)
            cell_border = Border(
                left=Side(style='thin', color='D9D9D9'),
                right=Side(style='thin', color='D9D9D9'),
                top=Side(style='thin', color='D9D9D9'),
                bottom=Side(style='thin', color='D9D9D9')
            )

            for cell in header_row:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = cell_border

            from groups.models import GroupEnrollment
            enrollments = GroupEnrollment.objects.filter(group=group, student__in=students)
            enr_map = {e.student_id: e for e in enrollments}

            for row_num, student in enumerate(students, 3):
                enr = enr_map.get(student.id)
                is_student_frozen = False
                frozen_date = None

                if enr and enr.is_frozen:
                    is_student_frozen = True
                    if enr.frozen_at:
                        frozen_date = enr.frozen_at.date()
                elif getattr(student, 'is_frozen', False):
                    is_student_frozen = True
                    if getattr(student, 'frozen_at', None):
                        frozen_date = student.frozen_at.date()

                name_val = f"{student.full_name} ❄️ (Muzlatilgan)" if is_student_frozen else student.full_name

                c1 = ws.cell(row=row_num, column=1, value=row_num - 2)
                c2 = ws.cell(row=row_num, column=2, value=name_val)
                c3 = ws.cell(row=row_num, column=3, value=student.phone or "-")
                c4 = ws.cell(row=row_num, column=4, value=student.parent_phone or "-")

                for c in [c1, c2, c3, c4]:
                    c.border = cell_border
                    if is_student_frozen:
                        c.fill = frozen_fill
                if is_student_frozen:
                    c2.font = frozen_font
                
                joined_at = student.joined_at.date() if student.joined_at else None
                
                for col_idx, d in enumerate(date_list, 5):
                    cell = ws.cell(row=row_num, column=col_idx)
                    cell.border = cell_border
                    
                    # Muzlatilgan kunlar mantiqi:
                    # Muzlatilgunga qadar bo'lgan kunlarda davomat saqlanadi,
                    # Muzlatilgandan keyingi kunlarda (d >= frozen_date) qor belgisi (❄️) tursin
                    is_day_frozen = False
                    if is_student_frozen:
                        if frozen_date:
                            if d >= frozen_date:
                                is_day_frozen = True
                        else:
                            # Agar aniq frozen_date topilmasa, mavjud davomat yozuvi bo'lmasa qor belgisi
                            if not att_data.get(student.id, {}).get(str(d)):
                                is_day_frozen = True

                    if is_day_frozen:
                        cell.value = "❄️"
                        cell.font = frozen_font
                        cell.fill = frozen_fill
                        cell.alignment = center_align
                        continue

                    att_record = att_data.get(student.id, {}).get(str(d))
                    
                    if att_record:
                        if not att_record.get('is_confirmed'):
                            cell.value = "!" # Davomat olinmagan (tasdiqlanmagan)
                            cell.font = Font(color="FF0000", bold=True)
                        elif att_record.get('is_present') is True:
                            cell.value = "+"
                            cell.font = Font(color="008000", bold=True)
                        else:
                            cell.value = "K" 
                            cell.font = Font(color="FF0000", bold=True)
                    elif joined_at and d < joined_at:
                        cell.value = "-"
                    else:
                        # Record yo'q bo'lsa va dars kuni bo'lsa (date_list dars kunlaridan iborat)
                        cell.value = "!"
                        cell.font = Font(color="FF0000", bold=True)

                    cell.alignment = center_align

            # Yuklab olish tarixini saqlash (xatolik bo'lsa ham report yuklanishi kerak)
            try:
                report_date = date_type(year, month, 1)
                ReportDownloadTrack.objects.get_or_create(
                    user=request.user,
                    report_type='attendance_monthly',
                    report_date=report_date
                )
            except Exception as e:
                logger.warning(
                    "ReportDownloadTrack saqlashda xatolik: user=%s, group=%s, month=%s-%s — %s",
                    request.user, group.id, year, month, e
                )

            from urllib.parse import quote
            # Format filename safely to avoid UnicodeEncodeError in Content-Disposition header
            raw_filename = f"davomat_{group.name}_{year}_{month}.xlsx"
            # Remove characters that can break the HTTP header or represent directory traversal
            safe_filename = raw_filename.replace('\n', '').replace('\r', '').replace('"', '_').replace('/', '_').replace('\\', '_')
            encoded_filename = quote(safe_filename)
            
            response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            response['Content-Disposition'] = f'attachment; filename="{encoded_filename}"; filename*=utf-8\'\'{encoded_filename}'
            try:
                wb.save(response)
            except Exception as e:
                logger.exception("Excel faylni yaratishda xatolik: group=%s, month=%s-%s", group.id, year, month)
                return Response(
                    {"detail": f"Excel faylni yaratishda xatolik: {str(e)}"},
                    status=status.HTTP_500_INTERNAL_SERVER_ERROR
                )
            return response

        # JSON Response
        report_data = []
        from groups.models import GroupEnrollment
        json_enrollments = GroupEnrollment.objects.filter(group=group, student__in=students)
        json_enr_map = {e.student_id: e for e in json_enrollments}

        for student in students:
            history = []
            joined_at = student.joined_at.date() if student.joined_at else None
            enr = json_enr_map.get(student.id)
            is_student_frozen = False
            frozen_date = None

            if enr and enr.is_frozen:
                is_student_frozen = True
                if enr.frozen_at:
                    frozen_date = enr.frozen_at.date()
            elif getattr(student, 'is_frozen', False):
                is_student_frozen = True
                if getattr(student, 'frozen_at', None):
                    frozen_date = student.frozen_at.date()

            for d in date_list:
                date_str = str(d)
                att_record = att_data.get(student.id, {}).get(date_str)
                
                is_day_frozen = False
                if is_student_frozen:
                    if frozen_date:
                        if d >= frozen_date:
                            is_day_frozen = True
                    else:
                        if not att_record:
                            is_day_frozen = True

                if is_day_frozen:
                    att_status = "frozen"
                elif joined_at and d < joined_at:
                    att_status = "not_joined"
                elif att_record:
                    if not att_record.get('is_confirmed'):
                        att_status = "not_taken"
                    elif att_record.get('is_present'):
                        att_status = "present"
                    else:
                        att_status = "absent"
                else:
                    # Dars kuni, lekin rekord yo'q
                    if d < timezone.localdate():
                        att_status = "not_taken"
                    else:
                        att_status = "none"
                
                history.append({"date": date_str, "status": att_status})
            
            report_data.append({
                "student_id": student.id,
                "student_name": student.full_name,
                "is_frozen": is_student_frozen,
                "frozen_at": frozen_date.isoformat() if frozen_date else None,
                "attendance": history
            })

        return Response({
            "group_name": group.name,
            "period": f"{year}-{month:02d}",
            "data": report_data
        })

    @action(detail=True, methods=['patch'], url_path='edit-past', permission_classes=[IsAuthenticated])
    def edit_past(self, request, pk=None):
        """
        Oxirgi 3 ta o'tilgan darsdan birortasining davomatini tahrirlash.
        """
        attendance = self.get_object()
        serializer = AttendanceEditSerializer(data=request.data)
        
        if not serializer.is_valid():
            return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)
            
        # Ruxsat tekshirish: Admin, super_admin yoki ushbu guruhning mentori
        is_admin = request.user.role in ['admin', 'super_admin']
        is_group_mentor = request.user.role == 'mentor' and self._user_can_access_group(request.user, attendance.group)
        if not (is_admin or is_group_mentor):
            return Response(
                {"detail": "Sizda ushbu guruh davomatini tahrirlash huquqi yo'q."},
                status=status.HTTP_403_FORBIDDEN
            )
            
        try:
            from rest_framework.exceptions import ValidationError
            updated_attendance = edit_past_attendance(
                attendance_id=attendance.id,
                user=request.user,
                new_status=serializer.validated_data['new_status'],
                reason=serializer.validated_data['reason']
            )
            return Response(
                AttendanceSerializer(updated_attendance).data, 
                status=status.HTTP_200_OK
            )
        except ValidationError as e:
            return Response({"detail": str(e)}, status=status.HTTP_400_BAD_REQUEST)
        except Exception as e:
            logger.exception("Davomat tahrirlashda xatolik: %s", str(e))
            return Response(
                {"detail": "Tizim xatoligi yuz berdi. Iltimos adminga murojaat qiling."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )

    @action(detail=False, methods=['get'], url_path='student-monthly', permission_classes=[IsAuthenticated])
    def student_monthly(self, request):
        """
        O'quvchining 1 oylik davomatini GitHub-style heatmap uchun agregatsiya qilish.
        Faqat so'rov yuborilganda asinxron (lazy) olinadi.
        Moliya ma'lumotlariga mutlaqo daxl qilmaydi (Read-Only).
        """
        from groups.models import Student, GroupEnrollment
        from calendar import monthrange
        from datetime import date as date_type

        student_id = request.query_params.get('student_id')
        if not student_id:
            return Response({"detail": "student_id parametri majburiy"}, status=status.HTTP_400_BAD_REQUEST)

        try:
            student_id_int = int(student_id)
        except (ValueError, TypeError):
            return Response({"detail": "student_id butun son bo'lishi kerak"}, status=status.HTTP_400_BAD_REQUEST)

        student = get_object_or_404(Student, id=student_id_int)

        # Ruxsatlarni tekshirish (Access Control)
        user = request.user
        if user.role == 'admin':
            allowed_branches = set()
            if user.branch_id:
                allowed_branches.add(user.branch_id)
            if hasattr(user, 'branch_accesses'):
                allowed_branches.update(user.branch_accesses.values_list('branch_id', flat=True))

            # Studentning barcha aloqador filiallarini to'playmiz:
            student_branches = set()
            if student.branch_id:
                student_branches.add(student.branch_id)
            if student.group_id and student.group and student.group.branch_id:
                student_branches.add(student.group.branch_id)

            active_enrollment_branches = set(
                student.enrollments.filter(is_active=True, group__branch__isnull=False)
                .values_list('group__branch_id', flat=True)
            )
            student_branches.update(active_enrollment_branches)

            if not student_branches:
                any_branches = set(
                    student.enrollments.filter(group__branch__isnull=False)
                    .values_list('group__branch_id', flat=True)
                )
                student_branches.update(any_branches)

            # Agar student birorta filialga aloqador bo'lsa va admin ushbu filiallarning
            # birortasiga ham ruxsatga ega bo'lmasa, kirishni taqiqlaymiz
            if student_branches and not allowed_branches.intersection(student_branches):
                return Response({"detail": "Sizda ushbu o'quvchi ma'lumotlarini ko'rish huquqi yo'q"}, status=status.HTTP_403_FORBIDDEN)

            # Self-healing: agar student.group yoki student.branch nofaol/eskirgan bo'lsa
            if hasattr(student, 'sync_active_group_and_branch'):
                try:
                    student.sync_active_group_and_branch()
                except Exception:
                    pass

        elif user.role == 'mentor':
            # 1. Mentorga 'students' moduli bo'yicha maxsus ruxsat berilgan bo'lsa (staff_permissions)
            try:
                perms = user.staff_permissions.permissions or {}
                has_students_perm = bool(perms.get('students', False))
            except Exception:
                has_students_perm = False

            if has_students_perm:
                allowed_branches = set()
                if user.branch_id:
                    allowed_branches.add(user.branch_id)
                if hasattr(user, 'branch_accesses'):
                    allowed_branches.update(user.branch_accesses.values_list('branch_id', flat=True))

                student_branches = set()
                if student.branch_id:
                    student_branches.add(student.branch_id)
                student_branches.update(
                    student.enrollments.filter(is_active=True, group__branch__isnull=False)
                    .values_list('group__branch_id', flat=True)
                )
                if student_branches and not allowed_branches.intersection(student_branches):
                    return Response({"detail": "Sizda ushbu o'quvchi ma'lumotlarini ko'rish huquqi yo'q"}, status=status.HTTP_403_FORBIDDEN)
            else:
                # Asosiy mentor yoki qo'shimcha (additional) mentor bo'lgan guruhlar
                mentor_group_ids = set(
                    Group.objects.filter(
                        models.Q(mentor=user) | models.Q(additional_mentors__mentor=user)
                    ).values_list('id', flat=True)
                )

                student_group_ids = set(
                    student.enrollments.filter(is_active=True).values_list('group_id', flat=True)
                )
                if student.group_id:
                    student_group_ids.add(student.group_id)

                if not mentor_group_ids.intersection(student_group_ids):
                    return Response({"detail": "Sizda ushbu o'quvchi ma'lumotlarini ko'rish huquqi yo'q"}, status=status.HTTP_403_FORBIDDEN)

            if hasattr(student, 'sync_active_group_and_branch'):
                try:
                    student.sync_active_group_and_branch()
                except Exception:
                    pass

        today = timezone.localdate()
        try:
            year = int(request.query_params.get('year', today.year))
            month = int(request.query_params.get('month', today.month))
            if year < 2000 or year > 2100 or month < 1 or month > 12:
                year, month = today.year, today.month
            date_type(year, month, 1)
        except (ValueError, TypeError):
            year, month = today.year, today.month

        # Xavfsiz req_group_id tekshiruvi (null, undefined, bo'sh satr kabilardan himoyalangan)
        raw_group_id = request.query_params.get('group_id')
        req_group_id = None
        if raw_group_id and str(raw_group_id).strip().lower() not in ['null', 'undefined', 'none', '']:
            try:
                req_group_id = int(raw_group_id)
            except (ValueError, TypeError):
                req_group_id = None

        try:
            from homework_attends.services import compute_student_monthly_attendance
            result = compute_student_monthly_attendance(
                student=student,
                year=year,
                month=month,
                req_group_id=req_group_id,
                today=today
            )
            return Response(result)
        except Exception as e:
            logger.exception("student_monthly action xatoligi: %s", str(e))
            return Response(
                {"detail": "Davomat ma'lumotlarini yuklashda xatolik yuz berdi."},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR
            )
